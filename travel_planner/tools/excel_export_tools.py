import io

from google.genai import types
from openpyxl import Workbook
from openpyxl.styles import Font

from google.adk.tools.tool_context import ToolContext

_XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def export_trip_plan_to_excel(
    tool_context: ToolContext,
    destination: str,
    start_date: str,
    end_date: str,
    total_budget: float,
    total_cost: float,
    currency: str,
    flight_summary: str,
    daily_itinerary: list[dict],
    packing_list: list[str],
) -> dict:
    """Exports a finalized trip plan to a downloadable Excel workbook.

    Only call this after presenting a final plan to the user, and only when
    they ask to export/save/download it — this isn't something to do
    automatically every time. Fill every field from the plan you already
    presented; don't invent or re-derive numbers.

    Args:
        destination: City name, e.g. "Lisbon".
        start_date: Trip start date, YYYY-MM-DD.
        end_date: Trip end date, YYYY-MM-DD.
        total_budget: The user's stated total budget.
        total_cost: The plan's actual total cost.
        currency: 3-letter currency code.
        flight_summary: A short plain-text summary of the chosen flight
            (carrier, times, stops, price).
        daily_itinerary: One dict per day, e.g. {"day": "Day 1", "morning":
            "...", "afternoon": "...", "evening": "...", "estimated_cost": "..."}.
            Missing keys per day are shown blank rather than erroring.
        packing_list: Flat list of packing items, e.g. ["Rain jacket",
            "Comfortable walking shoes", ...].

    Returns:
        {"status": "success", "filename": <str>, "artifact_version": <int>}
    """
    workbook = Workbook()

    _build_overview_sheet(
        workbook.active,
        destination=destination,
        start_date=start_date,
        end_date=end_date,
        total_budget=total_budget,
        total_cost=total_cost,
        currency=currency,
        flight_summary=flight_summary,
    )
    _build_itinerary_sheet(workbook.create_sheet("Itinerary"), daily_itinerary)
    _build_packing_sheet(workbook.create_sheet("Packing List"), packing_list)

    buffer = io.BytesIO()
    workbook.save(buffer)

    filename = f"trip_plan_{destination.replace(' ', '_')}.xlsx"
    version = await tool_context.save_artifact(
        filename=filename,
        artifact=types.Part.from_bytes(data=buffer.getvalue(), mime_type=_XLSX_MIME_TYPE),
    )

    return {"status": "success", "filename": filename, "artifact_version": version}


def _build_overview_sheet(sheet, *, destination, start_date, end_date, total_budget, total_cost, currency, flight_summary):
    sheet.title = "Overview"
    bold = Font(bold=True)
    rows = [
        ("Destination", destination),
        ("Start date", start_date),
        ("End date", end_date),
        ("Total budget", f"{total_budget} {currency}"),
        ("Total cost", f"{total_cost} {currency}"),
        ("Flight", flight_summary),
    ]
    for row_index, (label, value) in enumerate(rows, start=1):
        sheet.cell(row=row_index, column=1, value=label).font = bold
        sheet.cell(row=row_index, column=2, value=value)
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 70


def _build_itinerary_sheet(sheet, daily_itinerary: list[dict]):
    headers = ["Day", "Morning", "Afternoon", "Evening", "Estimated Cost"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for day in daily_itinerary:
        sheet.append([
            day.get("day", ""),
            day.get("morning", ""),
            day.get("afternoon", ""),
            day.get("evening", ""),
            day.get("estimated_cost", ""),
        ])

    for column, width in zip("ABCDE", [10, 35, 35, 35, 16]):
        sheet.column_dimensions[column].width = width


def _build_packing_sheet(sheet, packing_list: list[str]):
    sheet.append(["Item"])
    sheet["A1"].font = Font(bold=True)
    for item in packing_list:
        sheet.append([item])
    sheet.column_dimensions["A"].width = 40
